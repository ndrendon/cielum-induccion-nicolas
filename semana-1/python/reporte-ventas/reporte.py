"""Limpia un CSV de ventas y genera un reporte en Excel con una hoja por análisis.

Uso:
    python reporte.py --entrada ventas.csv --salida reporte.xlsx
"""

import argparse
import logging
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

COLUMNAS_ESPERADAS = ["id_venta", "fecha", "cliente", "ciudad", "producto", "cantidad", "precio_unitario"]
COLUMNAS_OBLIGATORIAS = ["fecha", "ciudad", "producto", "cantidad", "precio_unitario"]
FORMATOS_FECHA = ["%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"]
NOMBRES_COLUMNAS = {
    "mes": "Mes",
    "ciudad": "Ciudad",
    "producto": "Producto",
    "cantidad_ventas": "Cantidad de ventas",
    "unidades": "Unidades vendidas",
    "total": "Total vendido",
}


def leer_ventas(ruta: Path) -> pd.DataFrame:
    logger.info("Leyendo ventas desde %s", ruta)
    ventas = pd.read_csv(ruta, dtype=str, encoding="utf-8-sig")
    faltantes = [columna for columna in COLUMNAS_ESPERADAS if columna not in ventas.columns]
    if faltantes:
        raise ValueError(f"Al archivo le faltan las columnas: {', '.join(faltantes)}")
    logger.info("Filas leídas: %d", len(ventas))
    return ventas


def convertir_fechas(fechas: pd.Series) -> pd.Series:
    resultado = pd.to_datetime(fechas, format=FORMATOS_FECHA[0], errors="coerce")
    for formato in FORMATOS_FECHA[1:]:
        resultado = resultado.fillna(pd.to_datetime(fechas, format=formato, errors="coerce"))
    return resultado


def limpiar_ventas(ventas: pd.DataFrame) -> pd.DataFrame:
    datos = ventas[COLUMNAS_ESPERADAS].copy()
    for columna in COLUMNAS_ESPERADAS:
        datos[columna] = datos[columna].str.strip()
    datos = datos.mask(datos == "")
    datos["ciudad"] = datos["ciudad"].str.title()

    filas_antes = len(datos)
    datos = datos.dropna(subset=COLUMNAS_OBLIGATORIAS)
    logger.info("Filas con datos vacíos eliminadas: %d", filas_antes - len(datos))

    fechas = convertir_fechas(datos["fecha"])
    invalidas = fechas.isna()
    corregidas = ~invalidas & (datos["fecha"] != fechas.dt.strftime("%Y-%m-%d"))
    logger.info("Fechas corregidas al formato AAAA-MM-DD: %d", corregidas.sum())
    if invalidas.any():
        logger.warning(
            "Filas con fecha inválida eliminadas: %d (%s)",
            invalidas.sum(),
            ", ".join(datos.loc[invalidas, "fecha"]),
        )
    datos = datos.assign(fecha=fechas).loc[~invalidas]

    datos = datos.assign(
        cantidad=pd.to_numeric(datos["cantidad"], errors="coerce"),
        precio_unitario=pd.to_numeric(datos["precio_unitario"], errors="coerce"),
    )
    numeros_invalidos = ~((datos["cantidad"] > 0) & (datos["precio_unitario"] > 0))
    if numeros_invalidos.any():
        logger.warning("Filas con cantidad o precio inválido eliminadas: %d", numeros_invalidos.sum())
    datos = datos.loc[~numeros_invalidos]

    filas_antes = len(datos)
    datos = datos.drop_duplicates()
    logger.info("Ventas duplicadas eliminadas: %d", filas_antes - len(datos))

    datos = datos.assign(cantidad=datos["cantidad"].astype("int64"))
    datos = datos.assign(total=datos["cantidad"] * datos["precio_unitario"]).reset_index(drop=True)
    logger.info("Ventas válidas después de limpiar: %d", len(datos))
    return datos


def total_por_mes(ventas: pd.DataFrame) -> pd.DataFrame:
    return (
        ventas.assign(mes=ventas["fecha"].dt.strftime("%Y-%m"))
        .groupby("mes", as_index=False)
        .agg(cantidad_ventas=("total", "size"), total=("total", "sum"))
    )


def total_por_ciudad(ventas: pd.DataFrame) -> pd.DataFrame:
    return (
        ventas.groupby("ciudad", as_index=False)
        .agg(cantidad_ventas=("total", "size"), total=("total", "sum"))
        .sort_values("total", ascending=False, ignore_index=True)
    )


def top_productos(ventas: pd.DataFrame, limite: int = 5) -> pd.DataFrame:
    return (
        ventas.groupby("producto", as_index=False)
        .agg(unidades=("cantidad", "sum"), total=("total", "sum"))
        .sort_values("total", ascending=False, ignore_index=True)
        .head(limite)
    )


def dar_formato(hoja: Worksheet) -> None:
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(fill_type="solid", start_color="1F4E78")
    hoja.freeze_panes = "A2"

    for columna in hoja.iter_cols():
        if columna[0].value == NOMBRES_COLUMNAS["total"]:
            for celda in columna[1:]:
                celda.number_format = '"$"#,##0'
        ancho = max(len(str(celda.value)) for celda in columna)
        hoja.column_dimensions[get_column_letter(columna[0].column)].width = ancho + 4


def exportar_excel(hojas: dict[str, pd.DataFrame], ruta: Path) -> None:
    with pd.ExcelWriter(ruta, engine="openpyxl") as escritor:
        for nombre_hoja, tabla in hojas.items():
            tabla.rename(columns=NOMBRES_COLUMNAS).to_excel(escritor, sheet_name=nombre_hoja, index=False)
            dar_formato(escritor.sheets[nombre_hoja])
            logger.info("Hoja '%s' creada con %d filas", nombre_hoja, len(tabla))
    logger.info("Reporte guardado en %s", ruta)


def crear_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Limpia un CSV de ventas y genera un reporte en Excel con una hoja por análisis.")
    parser.add_argument("--entrada", type=Path, required=True, help="ruta del CSV de ventas")
    parser.add_argument("--salida", type=Path, default=Path("reporte.xlsx"), help="ruta del Excel a generar (por defecto: reporte.xlsx)")
    return parser


def main(argumentos: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)-7s | %(message)s", datefmt="%H:%M:%S")
    args = crear_parser().parse_args(argumentos)

    try:
        ventas = limpiar_ventas(leer_ventas(args.entrada))
        hojas = {
            "Total por mes": total_por_mes(ventas),
            "Total por ciudad": total_por_ciudad(ventas),
            "Top 5 productos": top_productos(ventas),
        }
        exportar_excel(hojas, args.salida)
    except FileNotFoundError:
        logger.error("No se encontró el archivo de entrada: %s", args.entrada)
        return 1
    except PermissionError:
        logger.error("No se pudo escribir %s. Si está abierto en Excel, ciérralo e intenta de nuevo", args.salida)
        return 1
    except ValueError as error:
        logger.error("El archivo de entrada no es válido: %s", error)
        return 1

    return 0


if __name__ == "__main__":
    sys.exit(main())
