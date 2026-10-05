from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from reporte import leer_ventas, limpiar_ventas, main, top_productos, total_por_ciudad, total_por_mes

CSV_DE_PRUEBA = """id_venta,fecha,cliente,ciudad,producto,cantidad,precio_unitario
1,2026-01-10,Ana Gómez,Medellín,Teclado mecánico,2,180000
2,15/01/2026,Luis Pérez,  cali ,Mouse inalámbrico,1,65000
3,2026-02-30,Sofía Ruiz,Bogotá,Teclado mecánico,1,180000
4,2026-02-05,Mateo Díaz,,Mouse inalámbrico,3,65000
5,2026-02-07,Carlos Mora,Bogotá,Monitor 24 pulgadas,1,720000
5,2026-02-07,Carlos Mora,Bogotá,Monitor 24 pulgadas,1,720000
"""


@pytest.fixture
def ruta_csv(tmp_path: Path) -> Path:
    ruta = tmp_path / "ventas.csv"
    ruta.write_text(CSV_DE_PRUEBA, encoding="utf-8")
    return ruta


@pytest.fixture
def ventas_limpias(ruta_csv: Path) -> pd.DataFrame:
    return limpiar_ventas(leer_ventas(ruta_csv))


def test_limpiar_elimina_filas_con_datos_vacios(ventas_limpias: pd.DataFrame) -> None:
    assert "4" not in ventas_limpias["id_venta"].tolist()


def test_limpiar_corrige_fechas_mal_escritas_y_elimina_las_invalidas(ventas_limpias: pd.DataFrame) -> None:
    fechas = dict(zip(ventas_limpias["id_venta"], ventas_limpias["fecha"]))

    assert fechas["2"] == pd.Timestamp("2026-01-15")
    assert "3" not in fechas


def test_limpiar_elimina_duplicados_y_normaliza_ciudades(ventas_limpias: pd.DataFrame) -> None:
    assert ventas_limpias["id_venta"].tolist() == ["1", "2", "5"]
    assert ventas_limpias["ciudad"].tolist() == ["Medellín", "Cali", "Bogotá"]


def test_total_por_mes_suma_las_ventas_de_cada_mes(ventas_limpias: pd.DataFrame) -> None:
    resumen = total_por_mes(ventas_limpias)

    assert resumen["mes"].tolist() == ["2026-01", "2026-02"]
    assert resumen["cantidad_ventas"].tolist() == [2, 1]
    assert resumen["total"].tolist() == [425000, 720000]


def test_total_por_ciudad_ordena_de_mayor_a_menor(ventas_limpias: pd.DataFrame) -> None:
    resumen = total_por_ciudad(ventas_limpias)

    assert resumen["ciudad"].tolist() == ["Bogotá", "Medellín", "Cali"]
    assert resumen["total"].tolist() == [720000, 360000, 65000]


def test_top_productos_devuelve_solo_los_5_que_mas_venden() -> None:
    ventas = pd.DataFrame(
        {
            "producto": ["A", "B", "C", "D", "E", "F", "A"],
            "cantidad": [1, 1, 1, 1, 1, 1, 2],
            "total": [100, 600, 300, 450, 200, 50, 400],
        }
    )

    top = top_productos(ventas)

    assert top["producto"].tolist() == ["B", "A", "D", "C", "E"]
    assert top.loc[top["producto"] == "A", "unidades"].item() == 3


def test_main_genera_un_excel_con_una_hoja_por_analisis(ruta_csv: Path, tmp_path: Path) -> None:
    ruta_excel = tmp_path / "reporte.xlsx"

    codigo_salida = main(["--entrada", str(ruta_csv), "--salida", str(ruta_excel)])

    libro = load_workbook(ruta_excel)
    assert codigo_salida == 0
    assert libro.sheetnames == ["Total por mes", "Total por ciudad", "Top 5 productos"]
    encabezados = [celda.value for celda in libro["Top 5 productos"][1]]
    assert encabezados == ["Producto", "Unidades vendidas", "Total vendido"]


def test_main_devuelve_error_si_no_existe_el_archivo(tmp_path: Path) -> None:
    codigo_salida = main(["--entrada", str(tmp_path / "no_existe.csv"), "--salida", str(tmp_path / "reporte.xlsx")])

    assert codigo_salida == 1
