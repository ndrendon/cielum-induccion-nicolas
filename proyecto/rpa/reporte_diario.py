"""Bot de RPA del proyecto integrador (parte 7): reporte diario de pedidos.

Entra al frontend de React con un usuario de prueba, lee la tabla de pedidos
página por página y genera salida/reporte_diario.xlsx con los pedidos del día.

Cómo se corre (dentro de proyecto/rpa, con el entorno virtual activo):
    python reporte_diario.py                     pedidos de hoy
    python reporte_diario.py --fecha 2026-10-01  pedidos de otra fecha
    python reporte_diario.py --ver               muestra el navegador mientras trabaja

La dirección del frontend y el usuario van en el archivo .env (ver .env.example).
"""

import argparse
import logging
import os
import re
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from playwright.sync_api import Error as ErrorPlaywright
from playwright.sync_api import expect, sync_playwright

CARPETA = Path(__file__).resolve().parent
SALIDA = CARPETA / 'salida'
ERRORES = SALIDA / 'errores'
REPORTE = SALIDA / 'reporte_diario.xlsx'

logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)-7s %(message)s', datefmt='%H:%M:%S')
log = logging.getLogger('bot-reporte')


class ErrorDelBot(Exception):
    """Un problema que el bot reconoce, por ejemplo un usuario o una contraseña que no sirven."""


def leer_configuracion():
    """Lee la dirección del frontend y el usuario de prueba del archivo .env."""
    load_dotenv(CARPETA / '.env')
    url = os.getenv('FRONTEND_URL', 'http://localhost:5173').rstrip('/')
    usuario = os.getenv('RPA_USUARIO')
    clave = os.getenv('RPA_CLAVE')
    if not usuario or not clave:
        raise SystemExit('Faltan RPA_USUARIO o RPA_CLAVE en el archivo .env (copia .env.example como .env y llénalo)')
    return url, usuario, clave


def iniciar_sesion(pagina, url, usuario, clave):
    """Llena el formulario de login como lo haría una persona."""
    log.info('Entrando a %s/login con el usuario %s', url, usuario)
    pagina.goto(f'{url}/login')
    # Los campos se buscan por su etiqueta y el botón por su texto: así el bot no depende
    # de la estructura del HTML, que puede cambiar
    pagina.get_by_label('Correo').fill(usuario)
    pagina.get_by_label('Contraseña').fill(clave)
    pagina.get_by_role('button', name='Entrar').click()

    # Si el login sale bien, el frontend lleva a la lista de pedidos; si no, muestra un aviso de error
    titulo = pagina.get_by_role('heading', name='Pedidos', exact=True)
    aviso_error = pagina.get_by_role('alert')
    expect(titulo.or_(aviso_error)).to_be_visible()
    if aviso_error.is_visible():
        raise ErrorDelBot(
            f'el login no funcionó ({aviso_error.inner_text()}). Revisa RPA_USUARIO y RPA_CLAVE en el .env'
        )
    log.info('Sesión iniciada')


def contar(cantidad, singular, plural):
    """contar(1, 'pedido', 'pedidos') → '1 pedido'; contar(3, ...) → '3 pedidos'."""
    return f'{cantidad} {singular if cantidad == 1 else plural}'


def texto_a_pesos(texto):
    """'$ 135.000' → 135000. El frontend muestra los pesos sin decimales."""
    return int(re.sub(r'\D', '', texto))


def leer_pedidos_del_dia(pagina, url, fecha):
    """Recorre la tabla de pedidos página por página y se queda con los de la fecha.

    La tabla trae los pedidos del más nuevo al más viejo. Por eso, cuando aparece
    un pedido de una fecha anterior, ya no hay más pedidos de ese día y el bot para.
    """
    buscada = fecha.isoformat()
    pedidos = []
    numero_pagina = 1

    # Después del login el frontend ya está en la lista de pedidos. El bot no escribe direcciones
    # en el navegador: el frontend guarda la sesión en memoria y la perdería al recargar la página.
    # Por eso se mueve con los botones de la página, como una persona.

    while True:
        # Espera a que la página muestre la tabla, o el aviso de que no hay pedidos
        tabla = pagina.get_by_role('table')
        sin_pedidos = pagina.get_by_text('No hay pedidos con ese filtro.')
        expect(tabla.or_(sin_pedidos)).to_be_visible()
        if sin_pedidos.is_visible():
            break

        # Mientras carga la página siguiente, el frontend sigue mostrando la anterior.
        # Por eso se espera a que la paginación diga el número de página que se pidió.
        expect(pagina.get_by_text(re.compile(rf'^Página {numero_pagina} de'))).to_be_visible()

        hay_mas_viejos = False
        for fila in tabla.locator('tbody').get_by_role('row').all():
            numero, fecha_pedido, cliente, estado, total, _ = fila.get_by_role('cell').all_inner_texts()
            if fecha_pedido == buscada:
                pedidos.append(
                    {
                        'id_pedido': int(numero),
                        'fecha': fecha_pedido,
                        'cliente': cliente.strip(),
                        'estado': estado.strip(),
                        'total': texto_a_pesos(total),
                    }
                )
            elif fecha_pedido < buscada:
                hay_mas_viejos = True
                break

        siguiente = pagina.get_by_role('button', name='Siguiente')
        log.info(
            'Página %d leída: %s del %s hasta ahora', numero_pagina, contar(len(pedidos), 'pedido', 'pedidos'), buscada
        )
        if hay_mas_viejos or siguiente.is_disabled():
            break
        siguiente.click()
        numero_pagina += 1

    return pedidos


def guardar_reporte(pedidos, fecha, url):
    """Escribe el Excel: una hoja con los pedidos del día y otra con el resumen por estado."""
    SALIDA.mkdir(exist_ok=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Pedidos del día'

    negrita = Font(bold=True)
    encabezado = PatternFill('solid', fgColor='DDEBF7')
    linea = Side(style='thin', color='9BC2E6')
    pesos = '"$" #,##0'

    hoja['A1'] = f'Reporte diario de pedidos: {fecha.isoformat()}'
    hoja['A1'].font = Font(bold=True, size=14)
    hoja['A2'] = f'Generado por el bot de RPA el {datetime.now():%Y-%m-%d %H:%M} leyendo {url}'
    hoja['A2'].font = Font(italic=True, color='666666')

    columnas = ['Pedido', 'Fecha', 'Cliente', 'Estado', 'Total']
    hoja.append([])
    hoja.append(columnas)
    for celda in hoja[4]:
        celda.font = negrita
        celda.fill = encabezado
        celda.border = Border(bottom=linea)

    if pedidos:
        for pedido in sorted(pedidos, key=lambda p: p['id_pedido']):
            hoja.append([pedido['id_pedido'], pedido['fecha'], pedido['cliente'], pedido['estado'], pedido['total']])
        ultima = hoja.max_row
        hoja.append(['Total del día', None, contar(len(pedidos), 'pedido', 'pedidos'), None, f'=SUM(E5:E{ultima})'])
        for celda in hoja[hoja.max_row]:
            celda.font = negrita
            celda.border = Border(top=linea)
        for fila in hoja.iter_rows(min_row=5, min_col=5, max_col=5):
            fila[0].number_format = pesos
        hoja.auto_filter.ref = f'A4:E{ultima}'
    else:
        hoja.append([f'No hubo pedidos el {fecha.isoformat()}'])

    for letra, ancho in zip('ABCDE', [14, 13, 28, 13, 16], strict=True):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = 'A5'

    # Hoja 2: cuántos pedidos y cuánto dinero hubo por estado
    resumen = libro.create_sheet('Resumen por estado')
    resumen.append(['Estado', 'Pedidos', 'Total'])
    for celda in resumen[1]:
        celda.font = negrita
        celda.fill = encabezado
    for estado in ['pendiente', 'enviado', 'entregado', 'cancelado']:
        del_estado = [p for p in pedidos if p['estado'] == estado]
        resumen.append([estado, len(del_estado), sum(p['total'] for p in del_estado)])
    resumen.append(['Total', len(pedidos), sum(p['total'] for p in pedidos)])
    for celda in resumen[resumen.max_row]:
        celda.font = negrita
    for fila in resumen.iter_rows(min_row=2, min_col=3, max_col=3):
        fila[0].number_format = pesos
    for letra, ancho in zip('ABC', [14, 10, 16], strict=True):
        resumen.column_dimensions[letra].width = ancho
    for fila in resumen.iter_rows(min_row=2, min_col=2, max_col=2):
        fila[0].alignment = Alignment(horizontal='center')

    libro.save(REPORTE)
    return REPORTE


def main():
    argumentos = argparse.ArgumentParser(description='Genera el reporte diario de pedidos leyendo el frontend.')
    argumentos.add_argument(
        '--fecha',
        type=date.fromisoformat,
        default=date.today(),
        help='fecha de los pedidos, con el formato AAAA-MM-DD (por defecto, hoy)',
    )
    argumentos.add_argument('--ver', action='store_true', help='muestra el navegador mientras el bot trabaja')
    opciones = argumentos.parse_args()

    url, usuario, clave = leer_configuracion()
    marca = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')

    with sync_playwright() as playwright:
        navegador = playwright.chromium.launch(headless=not opciones.ver, slow_mo=300 if opciones.ver else 0)
        pagina = navegador.new_page(viewport={'width': 1280, 'height': 900})
        # Si un elemento no aparece en 15 s, el bot falla en lugar de quedarse esperando para siempre
        pagina.set_default_timeout(15_000)
        expect.set_options(timeout=15_000)
        try:
            iniciar_sesion(pagina, url, usuario, clave)
            pedidos = leer_pedidos_del_dia(pagina, url, opciones.fecha)
            ruta = guardar_reporte(pedidos, opciones.fecha, url)
            total = sum(p['total'] for p in pedidos)
            log.info(
                'Listo: %s del %s por $ %s. Reporte en %s',
                contar(len(pedidos), 'pedido', 'pedidos'),
                opciones.fecha.isoformat(),
                f'{total:,}'.replace(',', '.'),
                ruta,
            )
        except (ErrorDelBot, ErrorPlaywright, AssertionError) as error:
            # Captura de la pantalla en el momento del error, para ver qué encontró el bot
            ERRORES.mkdir(parents=True, exist_ok=True)
            captura = ERRORES / f'error_{marca}.png'
            pagina.screenshot(path=captura, full_page=True)
            log.error('El bot falló: %s', str(error).splitlines()[0])
            if 'ERR_CONNECTION_REFUSED' in str(error):
                log.error('¿Está prendido el frontend en %s? Revisa FRONTEND_URL en el .env', url)
            log.error('Captura de la pantalla en %s', captura)
            raise SystemExit(1) from None
        finally:
            navegador.close()


if __name__ == '__main__':
    main()
